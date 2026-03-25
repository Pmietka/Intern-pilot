"""
AI-powered resume tailoring using Claude API.
Rewrites resume bullets for a specific job and generates a .docx file.
"""
import logging
from datetime import datetime
from pathlib import Path
from typing import Optional

import yaml

from ..tracker.database import Database
from ..tracker.models import Job
from ..utils.llm import call_claude_json

logger = logging.getLogger(__name__)

OUTPUTS_DIR = Path(__file__).resolve().parents[2] / "outputs" / "resumes"

RESUME_TAILOR_PROMPT = """\
You are a finance career advisor at a top university. Rewrite this candidate's resume bullets to be maximally relevant for the target job.

RULES:
- NEVER fabricate experience, metrics, or skills the candidate does not have
- Reorder bullets to put the most relevant experience first
- Mirror keywords from the job description naturally (for ATS optimization)
- Quantify impact wherever the existing data supports it
- Keep each bullet to one line, starting with a strong action verb
- Do NOT use em dashes, en dashes, or hyphens as punctuation. Use commas or restructure sentences instead.
- The tone should be professional but not robotic

CANDIDATE RESUME FACTS:
{resume_facts_from_profile}

TARGET JOB:
Title: {job_title}
Company: {company}
Description: {job_description}

Return the rewritten resume content as JSON:
{{
  "summary": "<optional 1 to 2 sentence professional summary if appropriate>",
  "experience": [
    {{
      "title": "...",
      "company": "...",
      "location": "...",
      "dates": "...",
      "bullets": ["...", "..."]
    }}
  ],
  "skills_section": ["skill1", "skill2"],
  "reordered_coursework": ["most relevant course"]
}}
"""


def build_resume_facts(profile: dict) -> str:
    """Format profile data into a structured string for the prompt."""
    lines = []
    personal = profile.get("personal", {})
    lines.append(f"Name: {personal.get('first_name', '')} {personal.get('last_name', '')}")
    lines.append(f"Email: {personal.get('email', '')}")
    lines.append(f"LinkedIn: {personal.get('linkedin', '')}")
    lines.append("")

    edu = profile.get("education", [{}])[0]
    lines.append(f"Education: {edu.get('degree', '')} in {edu.get('major', '')} at {edu.get('institution', '')}")
    lines.append(f"GPA: {edu.get('gpa', 'N/A')}")
    lines.append(f"Expected Graduation: {edu.get('expected_graduation', '')}")
    honors = edu.get("honors", [])
    if honors:
        lines.append(f"Honors: {', '.join(honors)}")
    coursework = edu.get("relevant_coursework", [])
    if coursework:
        lines.append(f"Relevant Coursework: {', '.join(coursework)}")
    lines.append("")

    lines.append("EXPERIENCE:")
    for exp in profile.get("experience", []):
        lines.append(
            f"  {exp.get('title', '')} at {exp.get('company', '')} "
            f"({exp.get('start_date', '')} to {exp.get('end_date', '')})"
        )
        for bullet in exp.get("bullets", []):
            lines.append(f"    - {bullet}")
    lines.append("")

    skills = profile.get("skills", {})
    tech = skills.get("technical", [])
    soft = skills.get("soft", [])
    lines.append(f"Technical Skills: {', '.join(tech)}")
    lines.append(f"Soft Skills: {', '.join(soft)}")

    return "\n".join(lines)


def tailor_resume(
    job: Job,
    profile: dict,
    delay: float = 2.0,
) -> dict:
    """
    Use Claude to generate tailored resume content for a specific job.

    Returns:
        Dict with keys: summary, experience, skills_section, reordered_coursework.
    """
    resume_facts = build_resume_facts(profile)
    description = job.description or "No description available."
    if len(description) > 4000:
        description = description[:4000] + "...[truncated]"

    prompt = RESUME_TAILOR_PROMPT.format(
        resume_facts_from_profile=resume_facts,
        job_title=job.title,
        company=job.company,
        job_description=description,
    )

    result = call_claude_json(prompt, delay=delay)
    logger.info("Resume tailored for job %d (%s at %s)", job.id, job.title, job.company)
    return result


def generate_resume_docx(
    job: Job,
    profile: dict,
    tailored_content: dict,
) -> Path:
    """
    Generate a tailored .docx resume file.

    Returns:
        Path to the generated .docx file.
    """
    try:
        from docx import Document
        from docx.shared import Pt, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        logger.error("python-docx is not installed. Run: pip install python-docx")
        raise

    OUTPUTS_DIR.mkdir(parents=True, exist_ok=True)

    personal = profile.get("personal", {})
    edu = profile.get("education", [{}])[0]

    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    def add_heading(text: str, level: int = 1) -> None:
        para = doc.add_heading(text, level=level)
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT

    def add_section_title(text: str) -> None:
        para = doc.add_paragraph()
        run = para.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(10)
        para.paragraph_format.space_after = Pt(2)
        # Add a simple divider using underline styling
        border_para = doc.add_paragraph("_" * 80)
        border_para.paragraph_format.space_after = Pt(4)
        border_para.paragraph_format.space_before = Pt(0)

    # Header: Name
    name = f"{personal.get('first_name', '')} {personal.get('last_name', '')}"
    header_para = doc.add_paragraph()
    header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header_para.add_run(name)
    run.bold = True
    run.font.size = Pt(16)

    # Contact info
    contact_parts = [
        personal.get("email", ""),
        personal.get("phone", ""),
        personal.get("linkedin", ""),
        f"{personal.get('address', {}).get('city', '')}, {personal.get('address', {}).get('state', '')}",
    ]
    contact_str = " | ".join(p for p in contact_parts if p and "[FILL IN]" not in p)
    contact_para = doc.add_paragraph(contact_str)
    contact_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_para.paragraph_format.space_after = Pt(6)

    # Summary (if provided)
    summary = tailored_content.get("summary", "")
    if summary:
        add_section_title("Summary")
        doc.add_paragraph(summary)

    # Education
    add_section_title("Education")
    edu_para = doc.add_paragraph()
    run = edu_para.add_run(
        f"{edu.get('institution', '')} | {edu.get('degree', '')} in {edu.get('major', '')}"
    )
    run.bold = True
    dates_run = edu_para.add_run(
        f"  |  Expected {edu.get('expected_graduation', '')}"
    )
    if edu.get("gpa") and "[FILL IN]" not in str(edu.get("gpa", "")):
        doc.add_paragraph(f"GPA: {edu.get('gpa')}")
    honors = edu.get("honors", [])
    if honors:
        doc.add_paragraph(f"Honors: {', '.join(honors)}")
    coursework = tailored_content.get("reordered_coursework") or edu.get("relevant_coursework", [])
    if coursework:
        doc.add_paragraph(f"Relevant Coursework: {', '.join(coursework)}")

    # Experience
    add_section_title("Experience")
    for exp in tailored_content.get("experience", []):
        exp_para = doc.add_paragraph()
        title_run = exp_para.add_run(f"{exp.get('title', '')} | {exp.get('company', '')}")
        title_run.bold = True
        dates_str = f"  |  {exp.get('dates', '')}  |  {exp.get('location', '')}"
        exp_para.add_run(dates_str)
        for bullet in exp.get("bullets", []):
            bullet_para = doc.add_paragraph(style="List Bullet")
            bullet_para.add_run(bullet)
            bullet_para.paragraph_format.space_after = Pt(1)

    # Skills
    skills_list = tailored_content.get("skills_section", [])
    if not skills_list:
        skills = profile.get("skills", {})
        skills_list = skills.get("technical", []) + skills.get("soft", [])
    if skills_list:
        add_section_title("Skills")
        doc.add_paragraph(", ".join(skills_list))

    # Save
    safe_company = "".join(c if c.isalnum() or c in " _" else "" for c in job.company).strip().replace(" ", "_")
    safe_title = "".join(c if c.isalnum() or c in " _" else "" for c in job.title).strip().replace(" ", "_")[:40]
    filename = f"resume_{safe_company}_{safe_title}_{job.id}.docx"
    output_path = OUTPUTS_DIR / filename

    doc.save(str(output_path))
    logger.info("Resume saved to %s", output_path)
    return output_path


def generate_resume_for_job(
    job_id: int,
    db: Database,
    profile: dict,
    delay: float = 2.0,
    dry_run: bool = False,
) -> Optional[Path]:
    """Generate a tailored resume for a specific job ID."""
    job = db.get_job(job_id)
    if not job:
        raise ValueError(f"Job {job_id} not found.")

    tailored = tailor_resume(job, profile, delay=delay)

    if dry_run:
        logger.info("[DRY RUN] Would generate resume for job %d", job_id)
        return None

    path = generate_resume_docx(job, profile, tailored)
    db.update_application_paths(job_id, resume_path=str(path))
    db.update_job_status(job_id, "content_generated")
    return path
